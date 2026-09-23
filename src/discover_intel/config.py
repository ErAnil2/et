"""Source registry: dataclass + CSV loader + xlsx seeder."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote_plus

import openpyxl

CSV_FILES = ("sources_web.csv", "sources_gnews.csv", "sources_youtube.csv")

# Beat queries from PRD § 5. Order preserved so the CSV is stable.
BEAT_QUERIES = [
    "Federal Reserve", "mortgage rates", "housing market", "S&P 500",
    "student loans", "Social Security", "Medicare", "IRS refund",
    "layoffs", "AI jobs", "electric vehicles", "mansion", "net worth",
    "credit card debt", "cost of living",
]

GNEWS_SECTIONS = [
    ("Top Stories", "https://news.google.com/rss?hl=en-US&gl=US&ceid=US:en"),
    ("Business",    "https://news.google.com/rss/headlines/section/topic/BUSINESS?hl=en-US&gl=US&ceid=US:en"),
    ("Technology",  "https://news.google.com/rss/headlines/section/topic/TECHNOLOGY?hl=en-US&gl=US&ceid=US:en"),
    ("World",       "https://news.google.com/rss/headlines/section/topic/WORLD?hl=en-US&gl=US&ceid=US:en"),
]


@dataclass(frozen=True)
class Source:
    source_id: str
    kind: str          # web|gnews_site|gnews_query|gnews_section|youtube
    market: str        # 'US'
    name: str
    url: str
    host: str | None
    tier: str | None
    category: str | None
    enabled: int
    notes: str | None


def _row_to_source(row: dict[str, str]) -> Source:
    def s(v: str | None) -> str | None:
        return v.strip() if (v is not None and v.strip() != "") else None
    return Source(
        source_id=row["source_id"].strip(),
        kind=row["kind"].strip(),
        market=row["market"].strip(),
        name=row["name"].strip(),
        url=row["url"].strip(),
        host=s(row.get("host")),
        tier=s(row.get("tier")),
        category=s(row.get("category")),
        enabled=int(row.get("enabled") or "1"),
        notes=s(row.get("notes")),
    )


def load_sources(config_dir: Path) -> list[Source]:
    out: list[Source] = []
    for fname in CSV_FILES:
        path = config_dir / fname
        if not path.exists():
            continue
        with path.open(encoding="utf-8", newline="") as fh:
            for row in csv.DictReader(fh):
                out.append(_row_to_source(row))
    return out


def _gnews_site_url(host: str) -> str:
    q = f"site:{host} when:1d"
    return f"https://news.google.com/rss/search?q={quote_plus(q)}&hl=en-US&gl=US&ceid=US:en"


def _gnews_query_url(q: str) -> str:
    quoted = f'"{q}" when:1d'
    return f"https://news.google.com/rss/search?q={quote_plus(quoted)}&hl=en-US&gl=US&ceid=US:en"


def _write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    cols = ["source_id", "kind", "market", "name", "url",
            "host", "tier", "category", "enabled", "notes"]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: ("" if r.get(c) is None else r.get(c)) for c in cols})


def seed_sources_from_xlsx(xlsx_path: Path, config_dir: Path) -> dict[str, int]:
    """Regenerate the three sources_*.csv files from a USA Top Publishers.xlsx."""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)

    # --- Sheet1: Publishers ---
    s1 = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    rows1 = list(s1.iter_rows(values_only=True))
    header1 = [str(c or "").strip() for c in rows1[0]]
    ix = {name: header1.index(name) for name in header1}
    pub_col = ix.get("Publisher", 0)
    rss_col = ix.get("RSS Feed Url", 1)

    web_rows: list[dict[str, object]] = []
    site_rows: list[dict[str, object]] = []
    for row in rows1[1:]:
        if row is None or row[pub_col] is None:
            continue
        host = str(row[pub_col]).strip().lower()
        rss = row[rss_col]
        # Every publisher gets a gnews_site row
        site_rows.append({
            "source_id": f"gnews:site:{host}", "kind": "gnews_site",
            "market": "US", "name": f"gnews site:{host}",
            "url": _gnews_site_url(host), "host": host, "enabled": 1,
        })
        # Publishers with a native RSS feed URL also get a web row
        if rss and str(rss).strip():
            web_rows.append({
                "source_id": f"web:{host}", "kind": "web", "market": "US",
                "name": host, "url": str(rss).strip(),
                "host": host, "enabled": 1,
            })

    # --- Beat queries + sections ---
    query_rows = [
        {
            "source_id": f"gnews:query:{q.lower().replace(' ', '-')}",
            "kind": "gnews_query", "market": "US", "name": q,
            "url": _gnews_query_url(q), "enabled": 1, "notes": "beat",
        }
        for q in BEAT_QUERIES
    ]
    section_rows = [
        {
            "source_id": f"gnews:section:{name.lower().replace(' ', '-')}",
            "kind": "gnews_section", "market": "US", "name": name,
            "url": url, "enabled": 1, "notes": "section",
        }
        for name, url in GNEWS_SECTIONS
    ]

    # --- Sheet2: YouTube channels ---
    s2 = wb["Sheet2"] if "Sheet2" in wb.sheetnames else wb.worksheets[1]
    rows2 = list(s2.iter_rows(values_only=True))
    header2 = [str(c or "").strip() for c in rows2[0]]
    jx = {name: header2.index(name) for name in header2}
    yt_rows: list[dict[str, object]] = []
    for row in rows2[1:]:
        if row is None:
            continue
        ch_id = row[jx.get("channel_id", 8)]
        if not ch_id or not str(ch_id).strip().startswith("UC"):
            continue
        ch_id = str(ch_id).strip()
        name = str(row[jx.get("channel_as_listed", 0)] or "").strip() or ch_id
        cat = row[jx.get("category", 1)]
        tier = row[jx.get("tier", 2)]
        yt_rows.append({
            "source_id": f"yt:{ch_id}", "kind": "youtube", "market": "US",
            "name": name,
            "url": f"https://www.youtube.com/feeds/videos.xml?channel_id={ch_id}",
            "host": "youtube.com",
            "tier": (str(tier).strip() if tier else None),
            "category": (str(cat).strip() if cat else None),
            "enabled": 1,
        })

    # --- Write files ---
    _write_csv(config_dir / "sources_web.csv", web_rows)
    _write_csv(config_dir / "sources_gnews.csv", site_rows + query_rows + section_rows)
    _write_csv(config_dir / "sources_youtube.csv", yt_rows)

    return {
        "web": len(web_rows),
        "gnews": len(site_rows) + len(query_rows) + len(section_rows),
        "youtube": len(yt_rows),
    }
